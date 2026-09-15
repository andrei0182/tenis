from __future__ import annotations

import os

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import TennisMatch

_COLUMN_LABELS = {
    "tour": "Tur",
    "tournament": "Turneu",
    "surface": "Suprafață",
    "time": "Ora",
    "status": "Status",
    "is_doubles": "Dublu",
    "player1": "Jucător 1",
    "player2": "Jucător 2",
    "odds_1": "Cotă 1",
    "odds_2": "Cotă 2",
    "games_line": "Linie Game-uri",
    "games_over": "Cotă Over",
    "games_under": "Cotă Under",
    "p1_ranking": "Rank J1",
    "p1_form": "Formă J1",
    "p1_surface_win_pct": "% Victorii Suprafață J1",
    "p2_ranking": "Rank J2",
    "p2_form": "Formă J2",
    "p2_surface_win_pct": "% Victorii Suprafață J2",
    "h2h_p1_wins": "H2H Victorii J1",
    "h2h_p2_wins": "H2H Victorii J2",
    "h2h_last_meeting": "Ultimul H2H",
    "estimated_edge": "Analiză / Edge Estimat",
    "match_url": "Link Meci",
}

_ODDS_COLUMNS = {"odds_1", "odds_2", "games_over", "games_under"}


def matches_to_dataframe(matches: list[TennisMatch]) -> pd.DataFrame:
    return pd.DataFrame(m.to_flat_dict() for m in matches)


def _style_sheet(ws, df: pd.DataFrame, odds_cols: set[str]) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    for col_idx, col_key in enumerate(df.columns, start=1):
        label = _COLUMN_LABELS.get(col_key, col_key)
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        letter = get_column_letter(col_idx)
        sample_values = [str(v) if pd.notna(v) else "" for v in df[col_key].head(200).tolist()]
        max_len = max([len(label)] + [len(v) for v in sample_values])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)

        if col_key in odds_cols:
            for row_idx in range(2, len(df) + 2):
                ws.cell(row=row_idx, column=col_idx).number_format = "0.00"

    ws.freeze_panes = "A2"

    if len(df) > 0:
        last_col_letter = get_column_letter(len(df.columns))
        table_ref = f"A1:{last_col_letter}{len(df) + 1}"
        table = Table(displayName="TennisMatchesTable", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)


def save_to_excel(matches: list[TennisMatch], path: str) -> None:
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)

    df = matches_to_dataframe(matches)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Tenis", index=False, header=False, startrow=1)
        _style_sheet(writer.sheets["Tenis"], df, _ODDS_COLUMNS)
